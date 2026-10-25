from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from db import get_conn

REPO = "TS_GP_IR_REPOMR_D.ORR"

SHEETS = [
    ("Flags", """
        SELECT fiscal_year, period_end, rule_code, variant, metric_value,
               warn_threshold, critical_threshold, severity
        FROM dbo.vw_fiscal_flags
        ORDER BY fiscal_year DESC, period_end DESC, rule_code"""),
    ("Revenue pace", """
        SELECT fiscal_year, period_end, target_type, target_pkr_mn, actual_pkr_mn,
               prior_year_share_pct, pct_of_phased_target
        FROM dbo.vw_r1_pace
        ORDER BY fiscal_year DESC, target_type, period_end"""),
    ("Quarterly flows", """
        SELECT q.fiscal_year, q.period_end, l.line_label, q.line_code,
               q.cumulative_pkr_mn, q.quarter_flow_pkr_mn, q.pub_status
        FROM dbo.vw_fiscal_quarterly q
        JOIN dbo.dim_fiscal_line l ON l.line_code = q.line_code
        ORDER BY q.fiscal_year DESC, q.period_end DESC, l.section, l.line_label"""),
    ("Review flags", """
        SELECT fiscal_year, period_end, line_code, quarter_flow_pkr_mn,
               mean_other_quarters, ratio_to_others, review_class
        FROM dbo.vw_fiscal_flow_review
        WHERE needs_review = 1
        ORDER BY fiscal_year DESC, line_code, period_end"""),
    ("Repo rate", f"""
        SELECT o.obs_date, o.value
        FROM dbo.market_observation o
        WHERE o.series_code = '{REPO}'
          AND o.loaded_at = (SELECT MAX(loaded_at) FROM dbo.market_observation
                             WHERE series_code = o.series_code AND obs_date = o.obs_date)
        ORDER BY o.obs_date DESC"""),
    ("Data quality", """
        SELECT source_file, check_name, max_difference, tolerance, passed, checked_at
        FROM dbo.dq_check_result ORDER BY check_id DESC"""),
    ("Run log", """
        SELECT TOP 100 run_id, pipeline, started_at, status, rows_loaded, error_message
        FROM dbo.etl_run_log ORDER BY run_id DESC"""),
]

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(bold=True, color="FFFFFF")
SEV_FILL = {"WARN": PatternFill("solid", fgColor="FFE699"),
            "CRITICAL": PatternFill("solid", fgColor="F4B084")}
TWO_DP = {"metric_value", "pct_of_phased_target", "prior_year_share_pct", "ratio_to_others",
          "warn_threshold", "critical_threshold", "max_difference", "tolerance", "value"}


def clean(v):
    return float(v) if isinstance(v, Decimal) else v


def fetch(cur, sql):
    cur.execute(sql)
    cols = [c[0] for c in cur.description]
    return cols, [[clean(v) for v in r] for r in cur.fetchall()]


def add_table(wb, name, cols, rows):
    ws = wb.create_sheet(name)
    ws.append(cols)
    for r in rows:
        ws.append(r)
    for c in ws[1]:
        c.fill, c.font = HEADER_FILL, HEADER_FONT
    for i, name_ in enumerate(cols, start=1):
        letter = get_column_letter(i)
        longest = max([len(str(name_))] + [len(str(r[i - 1])) for r in rows[:300]])
        ws.column_dimensions[letter].width = min(longest + 3, 48)
        fmt = "#,##0" if ("pkr_mn" in name_ or name_ == "mean_other_quarters") else \
              "0.00" if name_ in TWO_DP else None
        for cell in ws[letter][1:]:
            if isinstance(cell.value, datetime):
                cell.number_format = "yyyy-mm-dd hh:mm"
            elif isinstance(cell.value, date):
                cell.number_format = "yyyy-mm-dd"
            elif fmt and isinstance(cell.value, (int, float)):
                cell.number_format = fmt
    if "severity" in cols:
        col = cols.index("severity") + 1
        for row in ws.iter_rows(min_row=2):
            fill = SEV_FILL.get(row[col - 1].value)
            if fill:
                for c in row:
                    c.fill = fill
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    return ws


def one(cur, sql):
    cur.execute(sql)
    return cur.fetchall()


def main():
    conn = get_conn()
    cur = conn.cursor()
    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"

    latest_fiscal = one(cur, "SELECT MAX(period_end) FROM dbo.fiscal_cumulative")[0][0]
    latest_repo = one(cur, f"""SELECT TOP 1 obs_date, value FROM dbo.market_observation
                               WHERE series_code = '{REPO}' ORDER BY obs_date DESC, loaded_at DESC""")[0]
    sev = {r[0]: r[1] for r in one(cur, "SELECT severity, COUNT(*) FROM dbo.vw_fiscal_flags GROUP BY severity")}
    pace = one(cur, """SELECT fiscal_year, period_end, target_type, pct_of_phased_target
                       FROM dbo.vw_r1_pace
                       WHERE period_end = (SELECT MAX(period_end) FROM dbo.vw_r1_pace)
                       ORDER BY target_type""")

    rows = [
        ("Pakistan fiscal liquidity monitor", None),
        ("Report generated", datetime.now().strftime("%Y-%m-%d %H:%M")),
        ("Latest fiscal period in database", latest_fiscal),
        ("Latest repo rate date", latest_repo[0]),
        ("Latest overnight repo rate (percent)", float(latest_repo[1])),
        ("Flags raised: WARN", sev.get("WARN", 0)),
        ("Flags raised: CRITICAL", sev.get("CRITICAL", 0)),
    ]
    for fy, pe, tt, pct in pace:
        rows.append((f"Federal tax vs {tt.lower()} target, FY{fy} to {pe}, percent of phased target", float(pct)))
    rows.append(("All fiscal figures are PROVISIONAL as published; see the Notes sheet before use.", None))
    for r in rows:
        summary.append(list(r))
    summary["A1"].font = Font(bold=True, size=14)
    summary.column_dimensions["A"].width = 88
    summary.column_dimensions["B"].width = 20
    for row in summary.iter_rows(min_row=2, max_col=2):
        row[0].alignment = Alignment(wrap_text=True, vertical="top")
        if isinstance(row[1].value, date):
            row[1].number_format = "yyyy-mm-dd"

    for name, sql in SHEETS:
        cols, data = fetch(cur, sql)
        add_table(wb, name, cols, data)

    thresholds = fetch(cur, "SELECT rule_code, description, direction, warn_threshold, critical_threshold FROM dbo.rule_def")[1]
    notes = wb.create_sheet("Notes")
    lines = [
        "UNITS: PKR million unless a column says otherwise. Fiscal figures are cumulative from 1 July, as published; quarterly flows are derived by subtraction.",
        "SOURCES: Ministry of Finance Fiscal Operations, Table 1 (consolidated federal and provincial, inter-governmental balances excluded); SBP EasyData (overnight repo rate); budget targets from the Annual Budget Statements (federal FBR taxes only).",
        "STATUS: the fiscal figures carry the publisher's PROVISIONAL label. A later final version is stored as new rows, not overwritten.",
        "TARGETS: R1 compares Table 1 federal tax with the federal FBR taxes target. SBP Bulletin figures agree to rounding, but this is a comparison of related, not formally identical, definitions.",
        "PHASING: the annual target is phased by the previous year's cumulative share of federal tax. This is a modelling assumption, not a published schedule.",
        "ORIGINAL vs REVISED: the revised estimate is set with part of the year observed, so it is not an independent yardstick. The BUDGET variant is the early-warning one.",
        "LIMITS: three fiscal years are loaded. Flags are screening signals, not forecasts. R2, R3 and R4 have not fired on real data and were tested on constructed scenarios only.",
        "",
        "RULE THRESHOLDS (stored in dbo.rule_def):",
    ]
    for line in lines:
        notes.append([line])
    for t in thresholds:
        notes.append([f"{t[0]}: {t[1]} | fires {t[2]} warn={t[3]}, critical={t[4]}"])
    notes.column_dimensions["A"].width = 150
    for row in notes.iter_rows():
        row[0].alignment = Alignment(wrap_text=True, vertical="top")

    out = Path("reports") / f"fiscal_report_{date.today().isoformat()}.xlsx"
    out.parent.mkdir(exist_ok=True)
    wb.save(out)
    conn.close()
    print("saved:", out)


if __name__ == "__main__":
    main()